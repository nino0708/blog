#!/usr/bin/env python3
"""見やすいExcelの見本（references/excel.md の規則を実装したもの）。

  python3 sample_table.py out.xlsx

- 「説明」「入力」「一覧」の3シート構成
- 入力セルは青字、計算は黒字、別シート参照は緑字
- 罫線は横線だけ、枠線（グリッド線）は非表示、見出しは濃紺に白の太字
- 数値は右揃え＋桁区切り、見出し行とタイトル列を固定、テーブル化、A4横・1ページ幅で印刷
"""
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo

FONT = "BIZ UDPGothic"
NAVY, GRAY_LINE, BAND = "232F3E", "D5DBE3", "F1F4F8"
INPUT_BLUE, LINK_GREEN = "0B5CAD", "1D7A31"

base = Font(name=FONT, size=11)
bold = Font(name=FONT, size=11, bold=True)
head_font = Font(name=FONT, size=11, bold=True, color="FFFFFF")
head_fill = PatternFill("solid", fgColor=NAVY)
thin = Side(style="thin", color=GRAY_LINE)
strong = Side(style="medium", color=NAVY)


def style_header(ws, row, first_col, last_col):
    for c in range(first_col, last_col + 1):
        cell = ws.cell(row=row, column=c)
        cell.font, cell.fill = head_font, head_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[row].height = 24


def body_rows(ws, first_row, last_row, first_col, last_col):
    for r in range(first_row, last_row + 1):
        ws.row_dimensions[r].height = 20
        for c in range(first_col, last_col + 1):
            cell = ws.cell(row=r, column=c)
            if cell.font == Font():          # 未設定なら本文フォント
                cell.font = base
            cell.border = Border(bottom=thin)
            cell.alignment = Alignment(vertical="center", horizontal=cell.alignment.horizontal)


def build(path):
    wb = Workbook()

    # 説明シート
    ws = wb.active
    ws.title = "説明"
    ws.sheet_view.showGridLines = False
    ws["B2"] = "月別 作業時間の集計（見本）"
    ws["B2"].font = Font(name=FONT, size=14, bold=True)
    notes = [
        ("目的", "作業の種類ごとに、月の作業時間を集計する"),
        ("更新日", "2026-10-01"),
        ("入力", "前提の数字（単価・目標）。青字のセルだけ変更してよい"),
        ("一覧", "作業ごとの時間と金額。金額は入力シートの単価から計算"),
        ("色の意味", "青字＝入力、黒字＝計算、緑字＝別シートの参照"),
    ]
    for i, (k, v) in enumerate(notes, start=4):
        ws.cell(row=i, column=2, value=k).font = bold
        ws.cell(row=i, column=3, value=v).font = base
        ws.row_dimensions[i].height = 20
    ws.column_dimensions["A"].width = 2
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 60

    # 入力シート
    wi = wb.create_sheet("入力")
    wi.sheet_view.showGridLines = False
    wi.append([])
    wi.append([None, "項目", "単位", "値"])
    style_header(wi, 2, 2, 4)
    for name, unit, val in [("1時間あたりの単価", "円", 4500), ("月の目標時間", "時間", 140)]:
        wi.append([None, name, unit, val])
    body_rows(wi, 3, 4, 2, 4)
    for r in (3, 4):
        wi.cell(row=r, column=4).font = Font(name=FONT, size=11, color=INPUT_BLUE)
        wi.cell(row=r, column=4).number_format = "#,##0"
        wi.cell(row=r, column=4).alignment = Alignment(horizontal="right", vertical="center")
    for col, w in zip("ABCD", (2, 22, 8, 12)):
        wi.column_dimensions[col].width = w

    # 一覧シート
    wl = wb.create_sheet("一覧")
    wl.sheet_view.showGridLines = False
    months = ["7月", "8月", "9月"]
    header = ["区分", "作業", "単位"] + months + ["合計", "金額（円）"]
    wl.append(header)
    style_header(wl, 1, 1, len(header))
    rows = [
        ("定例", "設計レビュー", "時間", 42, 38, 45),
        ("定例", "資料作成", "時間", 18, 20, 16),
        ("定例", "月次報告", "時間", 8, 8, 9),
        ("非定例", "障害対応", "時間", 12, 25, 30),
        ("非定例", "技術調査", "時間", 4, 0, 6),
    ]
    for i, (cat, task, unit, *vals) in enumerate(rows, start=2):
        wl.append([cat, task, unit, *vals])
        wl.cell(row=i, column=7, value=f"=SUM(D{i}:F{i})")
        price = wl.cell(row=i, column=8, value=f"=G{i}*入力!$D$3")
        price.font = Font(name=FONT, size=11, color=LINK_GREEN)
    last = len(rows) + 1
    total = last + 1
    wl.cell(row=total, column=2, value="合計")
    for c in range(4, 9):
        col = wl.cell(row=1, column=c).column_letter
        wl.cell(row=total, column=c, value=f"=SUBTOTAL(109,{col}2:{col}{last})")
    body_rows(wl, 2, total, 1, len(header))
    for r in range(2, total + 1):
        for c in range(4, 9):
            cell = wl.cell(row=r, column=c)
            cell.number_format = "#,##0"
            cell.alignment = Alignment(horizontal="right", vertical="center")
        wl.cell(row=r, column=3).alignment = Alignment(horizontal="center", vertical="center")
    for c in range(1, len(header) + 1):              # 合計行：上に太い線、太字
        cell = wl.cell(row=total, column=c)
        cell.border = Border(top=strong, bottom=thin)
        cell.font = Font(name=FONT, size=11, bold=True, color=cell.font.color)
    tab = Table(displayName="作業一覧", ref=f"A1:{wl.cell(row=last, column=len(header)).coordinate}")
    tab.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=False)
    wl.add_table(tab)
    for col, w in zip("ABCDEFGH", (9, 18, 7, 9, 9, 9, 10, 14)):
        wl.column_dimensions[col].width = w
    wl.freeze_panes = "D2"
    ps = wl.page_setup
    ps.orientation, ps.paperSize = "landscape", wl.PAPERSIZE_A4
    ps.fitToWidth, ps.fitToHeight = 1, 0
    wl.sheet_properties.pageSetUpPr.fitToPage = True
    wl.print_title_rows = "1:1"
    wl.oddFooter.center.text = "&P / &N"

    wb.active = wb.index(wl)
    wb.save(path)
    return path


if __name__ == "__main__":
    print(build(sys.argv[1] if len(sys.argv) > 1 else "sample_table.xlsx"))
